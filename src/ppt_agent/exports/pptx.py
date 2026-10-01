"""Image-only PowerPoint exporter."""

from io import BytesIO

from pptx import Presentation
from pptx.util import Inches


def export_pptx(pages: tuple[bytes, ...]) -> bytes:
    if not pages:
        raise ValueError("at least one page is required")
    deck = Presentation()
    deck.slide_width = Inches(13.333333)
    deck.slide_height = Inches(7.5)
    deck.core_properties.title = "PPT Agent Export"
    deck.core_properties.author = "PPT Agent"
    blank = deck.slide_layouts[6]
    for content in pages:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(BytesIO(content), 0, 0, width=deck.slide_width, height=deck.slide_height)
    output = BytesIO()
    deck.save(output)
    return output.getvalue()
