from io import BytesIO
from zipfile import ZipFile

from PIL import Image
from pypdf import PdfReader
from pptx import Presentation

from ppt_agent.exports.images import convert_pages
from ppt_agent.exports.pdf import export_pdf
from ppt_agent.exports.pptx import export_pptx
from ppt_agent.exports.zip import export_zip


def _pages() -> tuple[bytes, ...]:
    result = []
    for color in ("red", "blue"):
        buffer = BytesIO()
        Image.new("RGB", (1920, 1080), color).save(buffer, "PNG")
        result.append(buffer.getvalue())
    return tuple(result)


def test_image_and_zip_exports_have_safe_ordered_names() -> None:
    jpg = convert_pages(_pages(), "jpg")
    assert Image.open(BytesIO(jpg[0])).size == (1920, 1080)
    archive = export_zip(jpg, "jpg")
    with ZipFile(BytesIO(archive)) as package:
        assert package.namelist() == ["slide-001.jpg", "slide-002.jpg"]


def test_jpeg_quality_is_passed_to_the_real_encoder() -> None:
    page = _pages()[0]

    low_quality = convert_pages((page,), "jpg", jpeg_quality=1)[0]
    high_quality = convert_pages((page,), "jpg", jpeg_quality=100)[0]

    assert low_quality != high_quality
    assert len(low_quality) < len(high_quality)


def test_pdf_is_16_by_9_and_pptx_has_one_image_per_slide() -> None:
    pages = _pages()
    pdf = PdfReader(BytesIO(export_pdf(pages)))
    assert len(pdf.pages) == 2
    assert round(float(pdf.pages[0].mediabox.width) / float(pdf.pages[0].mediabox.height), 3) == round(16 / 9, 3)
    deck = Presentation(BytesIO(export_pptx(pages)))
    assert len(deck.slides) == 2
    assert all(len(slide.shapes) == 1 for slide in deck.slides)


def test_pdf_dpi_is_passed_to_the_real_encoder() -> None:
    page = _pages()[0]

    at_72_dpi = PdfReader(BytesIO(export_pdf((page,), pdf_dpi=72)))
    at_144_dpi = PdfReader(BytesIO(export_pdf((page,), pdf_dpi=144)))

    assert float(at_72_dpi.pages[0].mediabox.width) == 1920
    assert float(at_144_dpi.pages[0].mediabox.width) == 960
